"""
Regression: _update_image_column sheet'i komple replace ediyordu
(openpyxl + if_sheet_exists='replace'). Kullanıcı manuel formatlama
(şartlı biçim, renk, validation) eklediyse siliniyordu.

Yeni davranış: openpyxl ile cell-level edit — sadece B kolonu hücreleri
güncellenir, geri kalan format korunur.
"""

import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, str(Path(__file__).parent.parent))

from images import ImageMatcher, MatchResult


def _build_excel_with_formatting(path: Path):
    """Çıktı sheet'i, A kolonunda formatlamayla Excel oluştur"""
    df = pd.DataFrame([
        {
            "Kart Listesi": "Player1 Series (1/5)",
            "Görsel Dosyası": "",
            "player_name": "Player1",
            "series_name": "Series",
            "group": "",
            "denominator": 5,
            "is_signed": "Hayır",
        },
        {
            "Kart Listesi": "Player2 Series (1/5)",
            "Görsel Dosyası": "",
            "player_name": "Player2",
            "series_name": "Series",
            "group": "",
            "denominator": 5,
            "is_signed": "Hayır",
        },
    ])
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Çıktı", index=False)

    # Manuel formatlama: A2 hücresine kalın yazı + sarı arka plan
    wb = load_workbook(path)
    ws = wb["Çıktı"]
    ws["A2"].font = Font(bold=True)
    ws["A2"].fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    ws.column_dimensions["A"].width = 99  # Özel kolon genişliği
    wb.save(path)


class TestUpdateImageColumnPreservesFormat:
    """_update_image_column çağrısı sheet formatlamasını korumalı"""

    def test_user_formatting_survives(self, tmp_path):
        excel = tmp_path / "test.xlsx"
        _build_excel_with_formatting(excel)
        image_dir = tmp_path / "images"
        image_dir.mkdir()

        matcher = ImageMatcher(excel, image_dir)
        # Sahte match'ler ekle (gerçek match akışını çalıştırmadan B kolonu güncelleme testi)
        matcher.matches = [
            MatchResult(row_number=2, card_text="Player1 Series (1/5)",
                        status="found", matched_file="player1_series_1.jpg"),
            MatchResult(row_number=3, card_text="Player2 Series (1/5)",
                        status="missing"),
        ]
        matcher._update_image_column()

        # Format korundu mu?
        wb = load_workbook(excel)
        ws = wb["Çıktı"]

        assert ws["A2"].font.bold is True, (
            "A2 kalın yazı stili silinmiş — sheet replace ediliyor olabilir"
        )
        fill = ws["A2"].fill
        # Solid PatternFill ise FFFFFF00 (sarı) formatında olmalı
        assert fill.fgColor.rgb is not None
        assert "FFFF00" in str(fill.fgColor.rgb).upper(), (
            f"A2 arka plan rengi gitmiş. fgColor: {fill.fgColor.rgb}"
        )
        assert ws.column_dimensions["A"].width == 99, (
            "A kolon genişliği değiştirilmiş"
        )

    def test_b_column_values_written(self, tmp_path):
        excel = tmp_path / "test.xlsx"
        _build_excel_with_formatting(excel)
        image_dir = tmp_path / "images"
        image_dir.mkdir()

        matcher = ImageMatcher(excel, image_dir)
        matcher.matches = [
            MatchResult(row_number=2, card_text="P1", status="found",
                        matched_file="player1_match.jpg"),
            MatchResult(row_number=3, card_text="P2", status="missing"),
        ]
        matcher._update_image_column()

        wb = load_workbook(excel)
        ws = wb["Çıktı"]
        assert ws["B2"].value == "player1_match.jpg"
        # missing için boş yazılmalı (None veya "")
        assert ws["B3"].value in (None, "")

    def test_conflict_status_written(self, tmp_path):
        excel = tmp_path / "test.xlsx"
        _build_excel_with_formatting(excel)
        image_dir = tmp_path / "images"
        image_dir.mkdir()

        matcher = ImageMatcher(excel, image_dir)
        matcher.matches = [
            MatchResult(row_number=2, card_text="P1", status="conflict",
                        conflict_files=["a.jpg", "b.jpg"]),
            MatchResult(row_number=3, card_text="P2", status="missing"),
        ]
        matcher._update_image_column()

        wb = load_workbook(excel)
        ws = wb["Çıktı"]
        assert "CONFLICT" in str(ws["B2"].value)
        assert "a.jpg" in str(ws["B2"].value)
        assert "b.jpg" in str(ws["B2"].value)

    def test_other_columns_unchanged(self, tmp_path):
        excel = tmp_path / "test.xlsx"
        _build_excel_with_formatting(excel)
        image_dir = tmp_path / "images"
        image_dir.mkdir()

        matcher = ImageMatcher(excel, image_dir)
        matcher.matches = [
            MatchResult(row_number=2, card_text="P1", status="found",
                        matched_file="x.jpg"),
        ]
        matcher._update_image_column()

        wb = load_workbook(excel)
        ws = wb["Çıktı"]
        assert ws["A2"].value == "Player1 Series (1/5)"
        assert ws["C2"].value == "Player1"
        assert ws["D2"].value == "Series"
        assert ws["F2"].value == 5
        assert ws["G2"].value == "Hayır"
